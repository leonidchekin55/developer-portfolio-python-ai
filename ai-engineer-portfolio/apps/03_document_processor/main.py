import uuid
import zipfile
from pathlib import Path
from fastapi import BackgroundTasks, Depends, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse
from sqlalchemy.orm import Session
from sqlalchemy import update
from sqlalchemy.exc import OperationalError
from pypdf import PdfReader
from docx import Document as WordDocument
from shared.core import create_app,get_db,Document,settings,user_from_token,SessionLocal
app=create_app("03 · AI Document Processor","Bounded batch upload, background extraction, persistent status and safe retry")
ALLOWED={".pdf",".docx",".txt",".md"}; MAX_FILE_BYTES=10*1024*1024; MAX_BATCH_BYTES=50*1024*1024; MAX_FILES=20

def extract(path:Path)->str:
    suffix=path.suffix.lower()
    if suffix==".pdf":
        reader=PdfReader(str(path))
        if len(reader.pages)>250: raise ValueError("PDF exceeds 250-page limit")
        pages=[]; total=0
        for page in reader.pages:
            part=page.extract_text() or ""; total+=len(part)
            if total>200_000: raise ValueError("Extracted text exceeds 200,000-character limit")
            pages.append(part)
        text="\n".join(pages)
    elif suffix==".docx":
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist())>50*1024*1024:
                raise ValueError("DOCX expands beyond 50 MiB")
        document=WordDocument(str(path)); text="\n".join(paragraph.text for paragraph in document.paragraphs)
    else: text=path.read_text(encoding="utf-8",errors="replace")
    if len(text)>200_000: raise ValueError("Extracted text exceeds 200,000-character limit")
    return text
def process(document_id:str):
    db=SessionLocal()
    try:
        claimed=db.execute(update(Document).where(Document.id==document_id,Document.status.in_(["queued","failed"])).values(status="processing",processing_error=None))
        db.commit()
        if claimed.rowcount!=1: return
        d=db.get(Document,document_id)
        if not d: return
        path=Path(settings.upload_dir)/f"{document_id}_{d.filename}"
        text=extract(path)
        d.text=text; d.status="processed" if text.strip() else "failed"
        d.processing_error=None if text.strip() else "EmptyExtractedText"
        db.commit()
    except OperationalError:
        db.rollback(); raise
    except Exception as exc:
        db.rollback()
        db.execute(update(Document).where(Document.id==document_id,Document.status=="processing").values(status="failed",processing_error=type(exc).__name__))
        db.commit()
    finally: db.close()
def enqueue_or_schedule(document_id:str,background:BackgroundTasks)->str:
    if settings.celery_enabled:
        from importlib import import_module
        worker=import_module("apps.03_document_processor.worker")
        worker.enqueue_document(document_id)
        return "redis/celery"
    background.add_task(process,document_id)
    return "in-process background task"
@app.post("/batch",status_code=202)
async def batch(background:BackgroundTasks,files:list[UploadFile]=File(...),db:Session=Depends(get_db),owner:str=Depends(user_from_token)):
    if not files: raise HTTPException(400,"At least one file is required")
    if len(files)>MAX_FILES: raise HTTPException(413,f"Maximum {MAX_FILES} files per batch")
    pending=[]; rejected=[]; batch_bytes=0
    for f in files:
        filename=Path(f.filename or "").name
        suffix=Path(filename).suffix.lower()
        if not filename or suffix not in ALLOWED:
            rejected.append({"filename":filename or "unknown","reason":"unsupported file type"}); continue
        data=await f.read(MAX_FILE_BYTES+1)
        if len(data)>MAX_FILE_BYTES:
            rejected.append({"filename":filename,"reason":"file exceeds 10 MiB"}); continue
        batch_bytes+=len(data)
        if batch_bytes>MAX_BATCH_BYTES: raise HTTPException(413,"Total batch size exceeds 50 MiB")
        ident=str(uuid.uuid4()); pending.append((ident,filename,data))
    if not pending and rejected: raise HTTPException(415,"No supported files were uploaded")
    folder=Path(settings.upload_dir); folder.mkdir(parents=True,exist_ok=True); written=[]
    try:
        for ident,filename,data in pending:
            path=folder/f"{ident}_{filename}"; path.write_bytes(data); written.append(path)
            db.add(Document(id=ident,owner=owner,filename=filename,status="queued",text=""))
        db.commit()
    except Exception:
        db.rollback()
        for path in written:
            if path.exists(): path.unlink()
        raise HTTPException(500,"Could not persist uploaded documents")
    try:
        queue_name="in-process background task"
        for ident,_,_ in pending:
            queue_name=enqueue_or_schedule(ident,background)
    except Exception as exc:
        raise HTTPException(503,"Documents were stored, but at least one task could not be queued; inspect statuses and retry") from exc
    return {"documents":[{"id":ident,"status":"queued"} for ident,_,_ in pending],"count":len(pending),"rejected":rejected,"queue":queue_name}
@app.get("/documents")
def list_documents(db:Session=Depends(get_db),owner:str=Depends(user_from_token)):
    return [{"id":d.id,"filename":d.filename,"status":d.status,"error":d.processing_error,"characters":len(d.text)} for d in db.query(Document).filter(Document.owner==owner).order_by(Document.created_at.desc()).limit(100)]
@app.get("/documents/{document_id}")
def status(document_id:str,db:Session=Depends(get_db),owner:str=Depends(user_from_token)):
    d=db.get(Document,document_id)
    if not d or d.owner!=owner: raise HTTPException(404,"Document not found")
    return {"id":d.id,"filename":d.filename,"status":d.status,"error":d.processing_error,"characters":len(d.text),"extracted_preview":d.text[:500]}
@app.post("/documents/{document_id}/retry",status_code=202)
def retry(document_id:str,background:BackgroundTasks,db:Session=Depends(get_db),owner:str=Depends(user_from_token)):
    d=db.get(Document,document_id)
    if not d or d.owner!=owner: raise HTTPException(404,"Document not found")
    if d.status not in {"failed","queued"}: raise HTTPException(409,"Only failed or queued documents can be retried")
    if not (Path(settings.upload_dir)/f"{d.id}_{d.filename}").exists(): raise HTTPException(410,"Original upload is unavailable")
    d.status="queued"; d.processing_error=None; db.commit()
    try: queue_name=enqueue_or_schedule(d.id,background)
    except Exception as exc: raise HTTPException(503,"Document is queued in the database but task dispatch failed") from exc
    return {"id":d.id,"status":"queued","queue":queue_name}
@app.get("/demo",response_class=HTMLResponse,include_in_schema=False)
def demo():
    token_default='value="demo-token-change-me"' if settings.app_env.lower() not in {"production","prod"} else ""
    return """<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Document Processor demo</title>
<style>body{font:16px system-ui;max-width:760px;margin:3rem auto;padding:0 1rem;color:#20242a}label{display:block;margin:.8rem 0}input,button{font:inherit;padding:.55rem}button{cursor:pointer}pre{white-space:pre-wrap;background:#f3f5f7;padding:1rem;border-radius:8px}</style>
<h1>Обработка документов</h1><p>Демо-обработка PDF, DOCX, TXT и Markdown. Максимум 20 файлов, 10 MiB каждый.</p>
<form id="upload"><label>Токен владельца <input id="token" type="password" autocomplete="off" __TOKEN_DEFAULT__ placeholder="Введите токен администратора"></label><label>Файлы <input id="files" type="file" multiple accept=".pdf,.docx,.txt,.md" required></label><button>Загрузить и обработать</button></form><h2>Результат</h2><pre id="result">Выберите файлы, чтобы увидеть статусы.</pre>
<script>
const form=document.querySelector('#upload'),out=document.querySelector('#result');
form.addEventListener('submit',async e=>{e.preventDefault();const data=new FormData();for(const f of document.querySelector('#files').files)data.append('files',f);out.textContent='Загрузка…';try{const token=document.querySelector('#token').value;const r=await fetch('/03-document-processor/batch',{method:'POST',headers:{Authorization:'Bearer '+token},body:data});const j=await r.json();if(!r.ok)throw new Error(j.detail||r.statusText);const lines=[];for(const d of j.documents||[]){let state=d.status,error='';for(let i=0;i<240&&state==='queued';i++){await new Promise(ok=>setTimeout(ok,500));const s=await fetch('/03-document-processor/documents/'+encodeURIComponent(d.id),{headers:{Authorization:'Bearer '+token}});if(s.ok){const status=await s.json();state=status.status;error=status.error||'';}}lines.push(d.id+': '+state+(error?' ('+error+')':''));}for(const x of j.rejected||[])lines.push(x.filename+': отклонён ('+x.reason+')');out.textContent=lines.join('\n')||'Нет обработанных файлов.';}catch(err){out.textContent='Ошибка: '+err.message;}});
</script></html>""".replace("__TOKEN_DEFAULT__",token_default)
